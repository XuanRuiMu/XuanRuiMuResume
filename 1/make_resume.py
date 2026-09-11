# -*- coding: utf-8 -*-
"""生成新版求职简历（AI 原生全栈定位版）· 一页紧凑版"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT_HEX = "1F4E79"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x59, 0x59, 0x59)
DARK = RGBColor(0x26, 0x26, 0x26)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"
EA = "微软雅黑"
TAB_RIGHT = Cm(18.2)  # 21 - 1.4*2 = 18.2cm 版心宽


def set_run(r, text=None, size=9.5, bold=False, color=DARK):
    if text is not None:
        r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), EA)
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)


def para(doc, before=0, after=1.5, line=1.03):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    return p


def shade(par, hex_color):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def bottom_border(par, hex_color=ACCENT_HEX, sz=10):
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), hex_color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def section(doc, title):
    p = para(doc, before=3, after=1.5)
    set_run(p.add_run(" " + title), size=10.5, bold=True, color=WHITE)
    shade(p, ACCENT_HEX)


def title_line(doc, left, sub, date, before=2):
    p = para(doc, before=before, after=0.5)
    p.paragraph_format.tab_stops.add_tab_stop(TAB_RIGHT, WD_TAB_ALIGNMENT.RIGHT)
    set_run(p.add_run(left), size=10, bold=True)
    if sub:
        set_run(p.add_run("　｜　" + sub), size=9, color=GRAY)
    set_run(p.add_run("\t" + date), size=8.5, color=GRAY)


def bullet(doc, lead, rest, after=1):
    p = para(doc, after=after)
    pf = p.paragraph_format
    pf.left_indent = Cm(0.42)
    pf.first_line_indent = Cm(-0.42)
    set_run(p.add_run("• "), size=9, bold=True, color=ACCENT)
    if lead:
        set_run(p.add_run(lead + "："), size=9.5, bold=True)
    set_run(p.add_run(rest), size=9.5)


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(1.4)
sec.right_margin = Cm(1.4)
sec.top_margin = Cm(0.8)
sec.bottom_margin = Cm(0.7)

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(9.5)
st._element.rPr.rFonts.set(qn("w:eastAsia"), EA)

# ---------- 头部 ----------
p = para(doc, after=0.5)
p.paragraph_format.tab_stops.add_tab_stop(TAB_RIGHT, WD_TAB_ALIGNMENT.RIGHT)
set_run(p.add_run("于翔堃"), size=18, bold=True, color=ACCENT)
set_run(p.add_run("\t2026 届应届本科 · 意向城市：天津（优先）"), size=9, color=GRAY)

p = para(doc, after=1)
set_run(p.add_run("AI 原生全栈开发者 ｜ 用 AI Agent 独立交付 6 个可运行系统，全流程可现场演示"),
        size=10, bold=True)

p = para(doc, after=1)
set_run(p.add_run("13312137630（微信同号）　|　3062949899@qq.com　|　github.com/mo-faa　|　在线作品站：101.42.45.157"),
        size=8.5, color=GRAY)
bottom_border(p)

# ---------- 求职意向 ----------
section(doc, "求职意向")
p = para(doc, after=0.5)
set_run(p.add_run("AI 应用开发 · 低代码开发 · 软件实施与交付 · 技术支持工程师"), size=9.5, bold=True, color=ACCENT)
set_run(p.add_run("　｜　随时到岗"), size=9, color=GRAY)

# ---------- 教育背景 ----------
section(doc, "教育背景")
title_line(doc, "天津仁爱学院", "计算机科学与技术 · 本科", "2022.09 – 2026.07", before=0.5)
bullet(doc, "主修", "数据结构、操作系统、计算机网络、数据库原理、软件工程")
bullet(doc, "毕业设计", "《基于 SSM 框架的网上村委会业务办理系统》独立开发（Spring MVC；通知发布与角色权限模块）")

# ---------- 核心能力 ----------
section(doc, "核心能力 · AI 原生工作流")
p = para(doc, after=1.5)
set_run(p.add_run("以 Claude Code / TRAE / Codex 等 AI 编程 Agent 为主要生产方式：AI 负责起草实现，我负责任务拆解、架构设计、代码审查、测试验证与上线决策，独立完成 6 个系统的全流程交付，可现场演示、可追问任意实现细节。"),
        size=9, color=GRAY)
bullet(doc, "AI Agent 工程", "ReAct 智能体闭环、Tool / Function Calling、RAG 混合检索、MCP 协议、SSE 流式输出、限流缓存与多级降级")
bullet(doc, "后端与数据库", "Python · FastAPI（异步）、Java · Spring Boot 3、MySQL 8（建模 / 优化 / 26.7GB 生产实例运维）、Redis")
bullet(doc, "前端与交互", "TypeScript、React 18、Astro SSR、Vue 3、原生 Canvas 游戏引擎、GSAP 动效")
bullet(doc, "部署与质量", "腾讯云、Nginx、Docker / compose、Linux、Git、pytest / JUnit 5 自动化测试、项目文档与验收自测")

# ---------- 项目经历 ----------
section(doc, "项目经历（独立开发 · 可现场演示）")

title_line(doc, "智能体工坊 —— ReAct Agent 运行时与 RAG 知识库平台", "Python 全栈", "2026.06", before=1)
bullet(doc, None, "从零实现 ReAct（推理-行动-观察）循环运行时与 6 个自研工具，支持自主规划与思考链路回放；自研 MCP 网关 SSE 流式推送推理过程（非框架套壳）")
bullet(doc, None, "混合检索 RAG：BM25 + 向量检索 + RRF 融合提升长尾召回；pytest + pytest-asyncio 覆盖核心链路")

title_line(doc, "AI 简历匹配平台", "Java 后端", "2026.08")
bullet(doc, None, "Spring Boot 3.5 + MySQL + Redis：粘贴招聘 JD 即输出匹配度分析与改进建议；Docker 多阶段构建 + compose 一键编排")
bullet(doc, None, "多级优雅降级（LLM 异常 → 规则引擎；Redis 故障 → 本地缓存）+ JD 哈希缓存 + IP 限流；JUnit 5 覆盖三条关键路径")

title_line(doc, "个人简历网站 + NEON CYBER 游戏平台（已上线：101.42.45.157）", "前端全栈", "2026.05 – 至今")
bullet(doc, None, "腾讯云 Lighthouse 部署（Nginx + Node 独立服务）公网可访问；Astro + TS + React 孤岛架构，主题系统 + WCAG AA 无障碍自检")
bullet(doc, None, "原生 JS + Canvas 7 款自研小游戏（粒子系统、对象池、WebAudio 音效、存档成就）；内置 JD 匹配器与 DeepSeek 流式 AI 求职助手")

title_line(doc, "会解压的外星人 —— AI 情绪陪伴应用", "Vue 全栈", "2026.04")
bullet(doc, None, "Vue 3 + Vite + Pinia；Node.js Express + SQLite；JWT 鉴权、SSE 流式对话、情绪标签识别与多语言支持")

title_line(doc, "暮澜纪元 Minecraft 服务器集群 —— 生产环境运维实战", "", "2024.03 – 至今")
bullet(doc, None, "长期运维 26.7GB MySQL 生产实例：备份、跨版本迁移、故障恢复，多次处置线上数据异常；打通权限 / 登录 / 经济插件数据链路")

# ---------- 岗位匹配 · 自我评价 ----------
section(doc, "岗位匹配 · 自我评价")
bullet(doc, "端到端交付", "6 个系统覆盖前后端、数据库、部署运维全链路，一人完成；项目文档齐全（做什么 / 怎么跑 / 在线演示）")
bullet(doc, "健壮性优先", "真实处置过大模型 API 余额耗尽（HTTP 402）故障，因降级链路服务未中断")
bullet(doc, "快速上手新领域", "Astro、Three.js、Rive 等均即学即用，可借 AI 工作流快速交付新平台业务；可接受出差 / 驻场")

OUT = r"C:\Users\17475\Desktop\03-于翔堃-求职简历-2026.docx"
doc.save(OUT)
print("saved:", OUT)
