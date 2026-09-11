import io, math, sys, subprocess
from PIL import ImageFont
from docx import Document

PLUGIN = r"C:\Users\17475\.workbuddy\plugins\cache\workbuddy-builtin\tencent-docx\5.5.3-wb.37748631.g104760a2.h1a8f7c37fe76"
OUT = r"C:\Users\17475\Desktop\项目\简历网页\output\69090144-c2f6-4bdd-bd2f-026198881595"
HTML = OUT + r"\stage2\formatted-于翔堃求职简历.html"
DOCX = r"C:\Users\17475\Desktop\求职资料\03-于翔堃-求职简历-2026.docx"

sys.path.insert(0, PLUGIN + r"\skills\html-to-docx\scripts")
from html_to_docx import convert, ConvertOptions  # noqa: E402

html = io.open(HTML, encoding="utf-8").read()
opts = ConvertOptions(page_size="A4", orientation="portrait",
                      margin_top=1.8, margin_bottom=1.8,
                      margin_left=1.8, margin_right=1.8, output_path=DOCX)
r = convert(html, output_path=DOCX, options=opts)
print("convert success =", r.success, "| error =", r.error, "| warnings =", r.warnings)

# ---------- 页数估算 ----------
FP = "C:/Windows/Fonts/msyh.ttc"
S = 8
_cache = {}


def font(sz):
    k = int(round(sz * S))
    if k not in _cache:
        _cache[k] = ImageFont.truetype(FP, k)
    return _cache[k]


PPMM = 72 / 25.4
content_pt = (210.0 - 18.0 - 18.0) * PPMM
page_pt = (297.0 - 18.0 - 18.0) * PPMM

d = Document(DOCX)
total = 0.0
rows = []
for p in d.paragraphs:
    pf = p.paragraph_format
    indent = pf.left_indent.pt if pf.left_indent else 0.0
    avail = content_pt - indent
    sz = None
    for run in p.runs:
        if run.font.size:
            sz = run.font.size.pt
            break
    if sz is None:
        sz = 9.5
    lh = 1.25 if p.style.name in ("Heading 1", "Heading 2") else 1.42
    width = 0.0
    for run in p.runs:
        s = run.font.size.pt if run.font.size else sz
        f = font(s)
        for ch in run.text:
            width += f.getlength(ch)
    width /= S
    lines = max(1, math.ceil(width / avail - 1e-6)) if width > 0 else 1
    h = lines * sz * lh
    sb = pf.space_before.pt if pf.space_before else 0.0
    sa = pf.space_after.pt if pf.space_after else 0.0
    total += h + sb + sa
    rows.append((p.style.name, lines, round(h + sb + sa, 1), p.text[:20]))

print("非空白段落 =", sum(1 for p in d.paragraphs if p.text.strip()))
print("内容总高 %.0f pt ｜ 每页可用 %.0f pt" % (total, page_pt))
print("估算页数 %.2f" % (total / page_pt))
print("--- 最高的 5 段 ---")
for row in sorted(rows, key=lambda x: -x[2])[:5]:
    print(row)
